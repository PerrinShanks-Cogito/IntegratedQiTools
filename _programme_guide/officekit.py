# © Guy’s and St Thomas’ NHS Foundation Trust 2026. All rights reserved.
"""
Small helpers for filling in the programme templates without losing their
formatting. Used by build.py.

  Deck   – a PowerPoint template (.potx or .pptx), via python-pptx
  Book   – an Excel workbook, edited as XML so pictures, charts, drop-downs
           and conditional formatting all survive
  detemplate() – turn a .potx/.dotx into a .pptx/.docx
"""
import copy, io, re, zipfile
from datetime import date, datetime
from lxml import etree

A = "http://schemas.openxmlformats.org/drawingml/2006/main"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"


def q(ns, tag):
    return "{%s}%s" % (ns, tag)


def detemplate(src):
    """Bytes of a .pptx/.docx made from a .potx/.dotx (or the file as it is)."""
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zi, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zo:
        for it in zi.infolist():
            data = zi.read(it.filename)
            if it.filename == "[Content_Types].xml":
                data = (data.replace(b"presentationml.template.main+xml", b"presentationml.presentation.main+xml")
                            .replace(b"wordprocessingml.template.main+xml", b"wordprocessingml.document.main+xml"))
            zo.writestr(it.filename, data)
    out.seek(0)
    return out


# --------------------------------------------------------------------------
# PowerPoint
# --------------------------------------------------------------------------
RAG = {"R": "FF0000", "A": "FFC000", "G": "92D050", "B": "00B0F0"}


def _run_props(p):
    """A copy of the first run's properties in paragraph p, as <a:rPr>."""
    r = p.find(q(A, "r") + "/" + q(A, "rPr"))
    if r is None:
        r = p.find(q(A, "endParaRPr"))
    rpr = copy.deepcopy(r) if r is not None else etree.Element(q(A, "rPr"))
    rpr.tag = q(A, "rPr")
    return rpr


def set_paras(txBody, lines, keep=0, italic=False, size=None, color=None, bold=None, template=None):
    """Replace the paragraphs of a text body from index `keep` onwards with
    `lines`, copying the paragraph and run formatting of the first replaced
    paragraph (or of paragraph `template`). A line is a string, or a list of
    (text, {"b": True}) pieces for mixed bold and plain text."""
    paras = txBody.findall(q(A, "p"))
    tpl = paras[template] if template is not None else (paras[keep] if len(paras) > keep else paras[-1])
    ppr = tpl.find(q(A, "pPr"))
    rpr = _run_props(tpl)
    rpr.attrib.pop("dirty", None)
    rpr.attrib.pop("err", None)
    for p in paras[keep:]:
        txBody.remove(p)
    if isinstance(lines, str):
        lines = [lines]
    for line in lines:
        p = etree.SubElement(txBody, q(A, "p"))
        if ppr is not None:
            p.append(copy.deepcopy(ppr))
        pieces = line if isinstance(line, list) else [(line, {})]
        for text, opts in pieces:
            r = etree.SubElement(p, q(A, "r"))
            rp = copy.deepcopy(rpr)
            rp.set("lang", "en-GB")
            rp.set("i", "1" if italic else "0")
            b = opts.get("b", bold)
            if b is not None:
                rp.set("b", "1" if b else "0")
            if size or opts.get("sz"):
                rp.set("sz", str(int((opts.get("sz") or size) * 100)))
            if color or opts.get("color"):
                for old in rp.findall(q(A, "solidFill")):
                    rp.remove(old)
                fill = etree.Element(q(A, "solidFill"))
                etree.SubElement(fill, q(A, "srgbClr"), val=opts.get("color") or color)
                # solidFill comes after <a:ln> but before effects and fonts
                ln = rp.find(q(A, "ln"))
                rp.insert(list(rp).index(ln) + 1 if ln is not None else 0, fill)
            if rp.find(q(A, "latin")) is None:
                # the templates set Arial on each run; the theme font is Calibri
                after = [rp.find(q(A, x)) for x in ("ea", "cs", "sym", "hlinkClick", "hlinkMouseOver", "rtl", "extLst")]
                after = [x for x in after if x is not None]
                latin = etree.Element(q(A, "latin"), typeface="Arial")
                if after:
                    after[0].addprevious(latin)
                else:
                    rp.append(latin)
            r.append(rp)
            t = etree.SubElement(r, q(A, "t"))
            t.text = text
        if not pieces:
            etree.SubElement(p, q(A, "endParaRPr"))
    return txBody


def set_para_text(txBody, idx, text):
    """Put text into paragraph idx, in its first run's formatting."""
    p = txBody.findall(q(A, "p"))[idx]
    runs = p.findall(q(A, "r"))
    if not runs:
        set_paras(txBody, [text], keep=idx, template=idx)
        return
    runs[0].find(q(A, "t")).text = text
    for r in runs[1:]:
        p.remove(r)
    for f in p.findall(q(A, "fld")):
        p.remove(f)


def fill_cell(cell, lines, **kw):
    set_paras(cell._tc.txBody, lines, **kw)


def add_rows(table, n):
    """Append n copies of a table's last row."""
    tbl = table._tbl
    last = tbl.tr_lst[-1]
    for _ in range(n):
        last.addnext(copy.deepcopy(last))
        last = tbl.tr_lst[-1]


def cell_colour(cell, hex_or_rag):
    """Set a table cell's background, e.g. "G" for the green RAG colour."""
    hexv = RAG.get(hex_or_rag, hex_or_rag)
    tcPr = cell._tc.get_or_add_tcPr()
    for old in tcPr.findall(q(A, "solidFill")) + tcPr.findall(q(A, "noFill")):
        tcPr.remove(old)
    fill = etree.Element(q(A, "solidFill"))
    etree.SubElement(fill, q(A, "srgbClr"), val=hexv)
    # fills come after the borders (lnL..lnBlToTr) in <a:tcPr>
    borders = [c for c in tcPr if etree.QName(c).localname.startswith("ln")]
    tcPr.insert(list(tcPr).index(borders[-1]) + 1 if borders else 0, fill)


def shape_fill(shape, hexv):
    shape.fill.solid()
    from pptx.dml.color import RGBColor
    shape.fill.fore_color.rgb = RGBColor.from_string(hexv)


class Deck:
    def __init__(self, path):
        from pptx import Presentation
        self.prs = Presentation(detemplate(path))

    @property
    def slides(self):
        return list(self.prs.slides)

    def shape(self, slide_no, ident):
        """Shape on slide (1-based) by shape id (int) or name (str)."""
        for sh in _all_shapes(self.slides[slide_no - 1].shapes):
            if (isinstance(ident, int) and sh.shape_id == ident) or sh.name == ident:
                return sh
        raise KeyError((slide_no, ident))

    def table(self, slide_no, ident):
        return self.shape(slide_no, ident).table

    def text(self, slide_no, ident, lines, keep=0, **kw):
        set_paras(self.shape(slide_no, ident).text_frame._txBody, lines, keep=keep, **kw)

    def para(self, slide_no, ident, idx, text):
        set_para_text(self.shape(slide_no, ident).text_frame._txBody, idx, text)

    def replace(self, slide_no, ident, old, new):
        """Swap text inside the runs of a shape, keeping their formatting."""
        sh = self.shape(slide_no, ident)
        for t in sh.text_frame._txBody.iter(q(A, "t")):
            if old in (t.text or ""):
                t.text = t.text.replace(old, new)

    def delete_slides(self, *numbers):
        lst = self.prs.slides._sldIdLst
        ids = list(lst)
        for n in sorted(numbers, reverse=True):
            sld = ids[n - 1]
            self.prs.part.drop_rel(sld.rId)
            lst.remove(sld)

    def delete_shape(self, slide_no, ident):
        sh = self.shape(slide_no, ident)
        sh._element.getparent().remove(sh._element)

    def save(self, path):
        self.prs.core_properties.title = ""
        self.prs.save(path)


def _all_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6:
            yield from _all_shapes(sh.shapes)


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------
def col_num(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def col_letters(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def split_ref(ref):
    m = re.match(r"([A-Z]+)(\d+)$", ref)
    return m.group(1), int(m.group(2))


DATE_IDS = set(range(14, 23)) | {45, 46, 47}


class Book:
    def __init__(self, path):
        with zipfile.ZipFile(path) as z:
            self.order = [i.filename for i in z.infolist()]
            self.files = {n: z.read(n) for n in self.order}
        self.xml = {}
        wb = self._x("xl/workbook.xml")
        rels = self._x("xl/_rels/workbook.xml.rels")
        targets = {r.get("Id"): r.get("Target") for r in rels}
        self.sheets = {}
        for sh in wb.iter(q(S, "sheet")):
            t = targets[sh.get(q(R, "id"))]
            self.sheets[sh.get("name")] = "xl/" + t.lstrip("/").replace("xl/", "", 1) if not t.startswith("/") else t.lstrip("/")
        self._date_xf = {}

    def _x(self, name):
        if name not in self.xml:
            self.xml[name] = etree.fromstring(self.files[name])
        return self.xml[name]

    def _styles(self):
        return self._x("xl/styles.xml")

    def _is_date_style(self, s):
        st = self._styles()
        xfs = st.find(q(S, "cellXfs"))
        xf = xfs[int(s)]
        fid = int(xf.get("numFmtId", "0"))
        if fid in DATE_IDS:
            return True
        nf = st.find(q(S, "numFmts"))
        if nf is not None:
            for f in nf:
                if int(f.get("numFmtId")) == fid:
                    code = re.sub(r'"[^"]*"|\[[^\]]*\]', "", f.get("formatCode")).lower()
                    return "d" in code and "y" in code
        return False

    def _date_style(self, s):
        """Index of a cell style like s but showing a date (dd/mm/yyyy)."""
        if self._is_date_style(s):
            return s
        if s in self._date_xf:
            return self._date_xf[s]
        st = self._styles()
        nf = st.find(q(S, "numFmts"))
        if nf is None:
            nf = etree.Element(q(S, "numFmts"), count="0")
            st.insert(0, nf)
        fid = None
        for f in nf:
            if f.get("formatCode") == "dd/mm/yyyy":
                fid = f.get("numFmtId")
        if fid is None:
            fid = str(max([164] + [int(f.get("numFmtId")) + 1 for f in nf]))
            etree.SubElement(nf, q(S, "numFmt"), numFmtId=fid, formatCode="dd/mm/yyyy")
            nf.set("count", str(len(nf)))
        xfs = st.find(q(S, "cellXfs"))
        xf = copy.deepcopy(xfs[int(s)])
        xf.set("numFmtId", fid)
        xf.set("applyNumberFormat", "1")
        xfs.append(xf)
        xfs.set("count", str(len(xfs)))
        self._date_xf[s] = str(len(xfs) - 1)
        return self._date_xf[s]

    def _wrap_style(self, s):
        """Index of a cell style like s that wraps text and aligns to the top."""
        key = ("wrap", s)
        if key in self._date_xf:
            return self._date_xf[key]
        st = self._styles()
        xfs = st.find(q(S, "cellXfs"))
        xf = copy.deepcopy(xfs[int(s)])
        al = xf.find(q(S, "alignment"))
        if al is None:
            al = etree.SubElement(xf, q(S, "alignment"))
        al.set("wrapText", "1")
        al.set("vertical", "top")
        xf.set("applyAlignment", "1")
        xfs.append(xf)
        xfs.set("count", str(len(xfs)))
        self._date_xf[key] = str(len(xfs) - 1)
        return self._date_xf[key]

    def cell(self, sheet, ref):
        root = self._x(self.sheets[sheet])
        data = root.find(q(S, "sheetData"))
        col, row = split_ref(ref)
        rows = data.findall(q(S, "row"))
        r_el = None
        for i, r in enumerate(rows):
            n = int(r.get("r"))
            if n == row:
                r_el = r
                break
            if n > row:
                r_el = etree.Element(q(S, "row"), r=str(row))
                r.addprevious(r_el)
                break
        if r_el is None:
            r_el = etree.SubElement(data, q(S, "row"), r=str(row))
        cn = col_num(col)
        for c in r_el.findall(q(S, "c")):
            other = col_num(split_ref(c.get("r"))[0])
            if other == cn:
                return c
            if other > cn:
                new = etree.Element(q(S, "c"), r=ref)
                c.addprevious(new)
                return new
        return etree.SubElement(r_el, q(S, "c"), r=ref)

    def set(self, sheet, ref, value, wrap=False, style=None):
        c = self.cell(sheet, ref)
        for ch in list(c):
            c.remove(ch)
        c.attrib.pop("t", None)
        if style is not None:
            c.set("s", str(style))
        s = c.get("s", "0")
        if value is None or value == "":
            return
        if isinstance(value, str) and value.startswith("="):
            etree.SubElement(c, q(S, "f")).text = value[1:]
            return
        if isinstance(value, (date, datetime)):
            c.set("s", self._date_style(s))
            etree.SubElement(c, q(S, "v")).text = str((value - date(1899, 12, 30)).days)
            return
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            etree.SubElement(c, q(S, "v")).text = repr(value) if isinstance(value, float) else str(value)
            return
        if wrap:
            c.set("s", self._wrap_style(s))
        c.set("t", "inlineStr")
        is_ = etree.SubElement(c, q(S, "is"))
        t = etree.SubElement(is_, q(S, "t"))
        t.text = str(value)
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")

    def style_of(self, sheet, ref):
        return self.cell(sheet, ref).get("s", "0")

    def row(self, sheet, row, start_col, values, wrap_cols=()):
        cn = col_num(start_col)
        for i, v in enumerate(values):
            letters = col_letters(cn + i)
            self.set(sheet, f"{letters}{row}", v, wrap=letters in wrap_cols)

    def row_height(self, sheet, row, points):
        root = self._x(self.sheets[sheet])
        self.cell(sheet, f"A{row}")
        for r in root.find(q(S, "sheetData")).findall(q(S, "row")):
            if r.get("r") == str(row):
                r.set("ht", str(points))
                r.set("customHeight", "1")

    def active_sheet(self, name):
        wb = self._x("xl/workbook.xml")
        names = [s.get("name") for s in wb.iter(q(S, "sheet"))]
        bv = wb.find(q(S, "bookViews"))[0]
        bv.set("activeTab", str(names.index(name)))
        bv.set("firstSheet", "0")
        for n, path in self.sheets.items():
            sv = self._x(path).find(q(S, "sheetViews"))
            if sv is not None and len(sv):
                if n == name:
                    sv[0].set("tabSelected", "1")
                else:
                    sv[0].attrib.pop("tabSelected", None)

    def save(self, path):
        # Excel recalculates every formula on opening, so the totals match the
        # values written here; the calculation chain is rebuilt then too.
        wb = self._x("xl/workbook.xml")
        cp = wb.find(q(S, "calcPr"))
        if cp is None:
            cp = etree.SubElement(wb, q(S, "calcPr"))
        cp.set("fullCalcOnLoad", "1")
        drop = {"xl/calcChain.xml"}
        rels = self._x("xl/_rels/workbook.xml.rels")
        for r in list(rels):
            if r.get("Target", "").endswith("calcChain.xml"):
                rels.remove(r)
        ct = etree.fromstring(self.files["[Content_Types].xml"])
        for o in list(ct):
            if o.get("PartName") == "/xl/calcChain.xml":
                ct.remove(o)
        self.xml["[Content_Types].xml"] = ct
        # drop the template's cached formula results so nothing stale shows
        for part in self.sheets.values():
            root = self._x(part)
            for c in root.iter(q(S, "c")):
                if c.find(q(S, "f")) is not None:
                    for v in c.findall(q(S, "v")):
                        c.remove(v)
                    if c.get("t") in ("str", "e", "b", "n"):
                        c.attrib.pop("t")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zo:
            for n in self.order:
                if n in drop:
                    continue
                data = etree.tostring(self.xml[n], xml_declaration=True, encoding="UTF-8", standalone=True) if n in self.xml else self.files[n]
                zo.writestr(n, data)
