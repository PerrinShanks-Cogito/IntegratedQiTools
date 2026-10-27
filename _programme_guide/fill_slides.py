# © Guy’s and St Thomas’ NHS Foundation Trust 2026. All rights reserved.
"""Fills in the PowerPoint templates with the worked example (example.py)."""
import copy
from lxml import etree
from pptx.util import Emu, Inches, Pt
from pptx.enum.shapes import MSO_CONNECTOR
from pptx.dml.color import RGBColor

from officekit import Deck, fill_cell, cell_colour, add_rows, set_paras, q, A
from example import *


def _rows(table, rows, start=1, **kw):
    for r, vals in enumerate(rows, start):
        for c, v in enumerate(vals):
            if v is None:
                continue
            fill_cell(table.cell(r, c), v if isinstance(v, list) else [v], **kw)


def _label(cell, label, value, **kw):
    """A cell like "SRO: [Name and title]": keep the bold label, swap the rest."""
    fill_cell(cell, [[(label + " ", {"b": True}), (value, {"b": False})]], **kw)


# --------------------------------------------------------------------------
def scope(src, dst):
    d = Deck(src)
    t = d.table(1, 4)
    _label(t.cell(0, 3), "SRO:", who("sro"))
    _label(t.cell(1, 3), "Key contact:", who("pm"))
    _label(t.cell(2, 3), "Parent programme:", "Urgent and Emergency Care Programme")
    fill_cell(t.cell(1, 0), [
        PITCH,
        "In place: this scope, SRO sign-off (5 Jan), CITI support plan (in draft). To follow: PID, resource plan, comms plan.",
    ])
    fill_cell(t.cell(4, 0), [
        "In scope: discharges from the 8 adult medical wards (Wards 1 to 8), Monday to Friday first, then weekends; board rounds, TTOs, transport bookings, the discharge lounge and complex discharges with our local authority partners.",
        "Out of scope: surgical and paediatric wards, the Emergency Department, community bed capacity and changes to the patient transport contract.",
    ])
    _rows(t, [[
        "Fewer A&E patients waiting for a bed in the afternoon and evening. About 3,000 bed days a year released (about £1.2m of cost avoided). Patients home earlier in the day, with their medicines. Supports the Trust strategy for urgent care.",
        "SRO Dr Amara Okafor; Programme Manager James Holloway with 4 CITI improvement staff. Ward time: each ward manager 2 hours a week, consultant leads 1 session a fortnight. Resource plan to follow.",
        "£35k capital for lounge chairs and refurbishment (agreed). £95k a year for 2 pharmacy technicians, to be agreed through the business case in March. CITI time from the existing CITI budget.",
        "Stakeholder map drafted: ward teams, pharmacy, patient transport, the discharge team, local authority social work, patients and carers, and the Urgent and Emergency Care Programme Board.",
    ]], start=6)
    _rows(t, [[
        "Fortnightly Programme Leadership Meeting chaired by the SRO, reporting monthly to the Medicine Clinical Group Executive Committee and the UEC Programme Board. Escalations go to the SRO.",
        "Ward teams can free 2 hours a week for the work. Pharmacy technician posts are approved by March. The transport provider can take morning bookings at no extra cost.",
        "Risk: winter pressures pull ward staff away from the work. Risk: the business case for pharmacy technicians is delayed. Issue: no reliable data yet on TTO ready times.",
        "Constraint: the discharge lounge has 12 chairs and 2 beds. Dependent on the Digital team for an EDD field in the board round view, and on the local authority for daily social work input.",
    ]], start=8)
    t2 = d.table(1, 5)
    fill_cell(t2.cell(0, 0), [GROUP])
    fill_cell(t2.cell(1, 0), [DIRECTORATE])
    fill_cell(t2.cell(2, 0), ["Completed 12 January 2026"])
    d.save(dst)


# --------------------------------------------------------------------------
def support_plan(src, dst):
    d = Deck(src)
    d.text(1, "Title 1", f"{NAME}: Support Plan")
    d.text(1, "Subtitle 2", "January 2026")
    # Context: | text | (blank) | Request for support: | text
    body = d.shape(2, "Content Placeholder 6").text_frame._txBody
    paras = body.findall(q(A, "p"))

    def like(p, text):
        n = copy.deepcopy(p)
        rs = n.findall(q(A, "r"))
        rs[0].find(q(A, "t")).text = text
        for r in rs[1:]:
            n.remove(r)
        return n
    context = [
        "Only 16% of patients on our 8 medical wards go home before midday, against 33% in the best-performing trusts. Beds free up late, so A&E patients wait for a bed into the evening and the wards start each day full.",
        "Average length of stay is 7.4 days. A discharge audit in November 2025 found TTOs ready on average 4.2 hours after prescribing, and 1 in 3 patients waiting on the ward for transport.",
        "A 2024 pilot of criteria-led discharge on Ward 6 raised morning discharges to 27% but was not kept up after the pilot ended.",
    ]
    request = [
        "6 months of CITI improvement support (March to August 2026), after a mobilisation month in February, to redesign the discharge day on all 8 wards with ward teams, pharmacy, transport and the discharge team.",
        "Areas affected: the 8 medical wards, pharmacy, the discharge lounge, patient transport bookings and the discharge team.",
    ]
    new = [paras[0]] + [like(paras[1], t) for t in context] + [paras[2], paras[3]] + [like(paras[4], t) for t in request]
    for p in paras:
        body.remove(p)
    for p in new:
        body.append(p)

    d.text(3, "Content Placeholder 6", "This table outlines the proposed scope for resources between CITI and the Medicine Clinical Group.")
    t = d.table(3, 4)
    _rows(t, [
        ["Board rounds and criteria-led discharge", "Standard morning board round on every ward; expected discharge date set within 24 hours; criteria-led discharge for suitable patients.", "Board round standard and role cards; criteria-led discharge policy; training for 8 ward teams.", "Discharges before midday; EDD set within 24 hours."],
        ["Medicines to take out (TTOs)", "Write TTOs the day before discharge; pharmacy technicians on wards; track TTO ready times.", "TTO process redesign; 2 pharmacy technician posts; daily TTO tracker.", "TTO ready time, prescribed to ready."],
        ["Transport, lounge and complex discharges", "Morning transport bookings; use the discharge lounge from 8am; daily transfer of care huddle.", "Transport booking process; lounge operating standard; transfer of care hub huddle.", "Lounge use; length of stay; readmissions (balancing)."],
    ], size=11)

    # Slide 4: approach
    d.para(4, 7, 0, "We will work through four Plan, Do, Study, Act stages with the ward teams, starting on Wards 5 and 6 and spreading to all 8 wards.")
    for ident, lines in [
        (64, ["Set up the leadership meeting, workstream groups and ward champions; launch the programme with all 8 wards",
              "Agree the board round and discharge standards with ward teams",
              "Build the baseline dashboard: discharges before midday, length of stay, TTO times, lounge use",
              "Finalise plans and link to the Urgent and Emergency Care Programme and the transfer of care hub"]),
        (65, ["Test the morning board round and criteria-led discharge on Wards 5 and 6, then spread ward by ward",
              "Write role cards and standard operating procedures for board rounds, TTOs and the lounge",
              "Coach ward teams on the ward, including bite-size training and visual management boards",
              "Work with pharmacy, transport and the discharge team to track the effect on their work"]),
        (68, ["Track the measures weekly on SPC charts and collect the blockers ward teams raise",
              "Run structured problem solving on the main blockers (TTO delays, transport, social care)",
              "Share each ward's measures at its daily huddle and in the weekly highlight report",
              "Agree how the changes will be kept up: audits, leader visits and ward champions"]),
        (67, ["Finalise the standard operating procedures, role cards and process documents",
              "Build the standards into induction, ward training and the Flow Manual audit",
              "Evaluate the impact and experience, and agree next steps with the Clinical Group",
              "Recommend how to spread the standards to surgical wards"]),
    ]:
        d.text(4, ident, lines)
    for ident, txt in [(52, "Mobilise and engage"), (53, "Agree standards"), (54, "Baseline metrics & targets"),
                       (55, "Test on 2 wards"), (56, "Train & coach"), (57, "Engage partners"),
                       (61, "Track impact weekly"), (62, "Solve blockers"), (63, "Sustainability plan"),
                       (58, "Finalise standards"), (59, "Embed in training"), (60, "Evaluate & spread")]:
        d.para(4, ident, 0, txt)

    # Slide 5: principles
    d.para(5, 11, 0, "The partnership between CITI and the Medicine Clinical Group aims to deliver change collaboratively, as one extended management and improvement team. Our approach will be pragmatic and will need involvement from leaders at all levels of the Clinical Group.")

    # Slide 6: governance
    d.text(6, 96, "The Discharge Flow Programme will use the Clinical Group's existing meetings where possible, so all improvement work is overseen in one place.", color="000000")
    t = d.table(6, 58)
    _rows(t, [
        [None, "Programme Leadership Meeting (fortnightly, chair: SRO)", "Medicine Clinical Group Executive Committee (monthly)*", "UEC Programme Board (monthly)*"],
        [None, "Weekly measures dashboard (SPC charts)", "Weekly highlight reports from each workstream", "Monthly benefits review"],
        [None, ["Programme weekly huddle (chair: Programme Manager)", "Workstream user groups"], "Daily ward discharge huddles*", "Transfer of care meeting (twice weekly)*"],
    ], start=0)
    fill_cell(t.cell(2, 0), [[("Delivery", {"b": True})], [("managed by CITI and ward teams", {"b": True})]])

    # Slide 7: programme management
    d.text(7, 9, [
        [("SRO: ", {"b": True}), (who("sro"), {"b": False})],
        [("Project Manager: ", {"b": True}), (who("pm"), {"b": False})],
        [("Clinical Lead: ", {"b": True}), (who("clin"), {"b": False})],
        [("Operational Lead: ", {"b": True}), (who("ops"), {"b": False})],
    ])
    _rows(d.table(7, 4), color="000000", rows=[
        ["Winter pressures pull ward staff away from the work", "Start on 2 wards in March; keep board rounds to 30 minutes; agree protected time with matrons"],
        ["Pharmacy technician business case delayed", "Take the case to the March investment panel; interim bank cover agreed"],
        ["No reliable data on TTO ready times", "Analyst to build a TTO tracker from the pharmacy system by end of February"],
        ["Changes not kept up after support ends (as in the 2024 pilot)", "Ward champions, Flow Manual audit and handover to the Clinical Group from July"],
    ])
    _rows(d.table(7, 10), color="000000", rows=[
        ["EDD field in the board round view", "Digital: electronic patient record team"],
        ["Daily social work input to the transfer of care huddle", "Local authority hospital social work team"],
        ["Same day emergency care changes", "Urgent and Emergency Care Programme"],
        ["Morning transport capacity", "Patient transport contract (Procurement)"],
    ])

    # Slide 8: benefits
    roi = d.table(8, 2)
    vals = [
        ["£35k", "£0", "£0", "£0", "£0", "£35k"],
        ["£0", "£0", "£0", "£0", "£0", "£0"],
        ["-£35k", "£0", "£0", "£0", "£0", "-£35k"],
        ["£180k", "£95k", "£95k", "£95k", "£95k", "£560k"],
        ["£600k", "£1.2m", "£1.2m", "£1.2m", "£1.2m", "£5.4m"],
        ["£420k", "£1.1m", "£1.1m", "£1.1m", "£1.1m", "£4.8m"],
    ]
    for r, row in enumerate(vals, 1):
        for c, v in enumerate(row, 2):
            fill_cell(roi.cell(r, c), [v])
    _rows(d.table(8, 3), [[k[0], k[1], k[2], "Sep 2026"] for k in KPI[:4]], color="000000")

    # Slide 9: resources
    d.para(9, 7, 0, "To resource these priority areas, we propose the following WTE from 2 February to 30 September 2026. This will let us support all 8 medical wards. Other initiatives will need further resources.")
    t = d.table(9, 5)
    _rows(t, [
        ["Programme leadership and PMO", "SRO 0.1 WTE; Accountable Executive 0.1 WTE", "Programme Director 0.2 WTE; Programme Manager 1.0 WTE; PMO 0.5 WTE"],
        ["Board rounds and criteria-led discharge", ["Board rounds: clinical lead 1 session a week; matron 0.2 WTE; ward managers 2 hours a week each", "TTOs: lead pharmacist 0.2 WTE; 2 pharmacy technicians (new posts)"], ["Board rounds: Improvement Manager 1.0 WTE", "TTOs: Improvement Manager 0.6 WTE"]],
        ["Medicines to take out (TTOs)", None, None],
        ["Transport, lounge and complex discharges", "Flow manager 0.2 WTE; discharge team lead 0.2 WTE; social work input to the daily huddle", "Project Manager 0.6 WTE; Improvement Manager 0.6 WTE"],
        ["Data and measurement", "Information team: data extracts", "Senior Analyst 0.4 WTE"],
    ])

    # Slide 10: next steps
    body = d.shape(10, 7).text_frame._txBody
    paras = body.findall(q(A, "p"))
    from officekit import set_para_text
    set_para_text(body, 0, "Confirm your approval of this Support Plan, including any amendments, by 23 January")
    set_para_text(body, 1, "Set up the kick-off meeting (the first in 2 weeks' time, on 2 February)")
    set_para_text(body, 2, "Identify members of the working group for:")
    for p in paras:
        for t in p.iter(q(A, "t")):
            if "Input forums" in (t.text or ""):
                t.text = "Medicine senior leadership team, ward managers' meeting, pharmacy team meeting, transfer of care meeting"
                rpr = t.getparent().find(q(A, "rPr"))
                for f in rpr.findall(q(A, "solidFill")):
                    for c in f:
                        f.remove(c)
                    from lxml import etree
                    etree.SubElement(f, q(A, "srgbClr"), val="000000")
    t = d.table(10, 3)
    _rows(t, [
        [[[("Board rounds group", {"b": True})], "Weekly, Tuesday 2pm"], "Dr Sofia Marquez, Joy Adebayo, 8 ward managers, Ellie Brooks, Ben Carter"],
        [[[("TTO group", {"b": True})], "Fortnightly, Thursday 10am"], "Leila Haddad, Arjun Mehta, 2 junior doctor reps, Hannah Lee"],
        [[[("Transport and discharge group", {"b": True})], "Fortnightly, Monday 1pm"], "Kwame Asante, Fiona Gallagher, Martin Doyle, Mark Ellis, Chloe Evans, Sam Patel"],
    ], start=0)
    d.save(dst)


# --------------------------------------------------------------------------
def mobilisation(src, dst):
    d = Deck(src)
    d.text(1, "Title 1", f"{NAME} Mobilisation Document")
    d.text(1, "Subtitle 2", long_date(date(2026, 2, 2)))
    d.text(3, 3, "Ward 6 photo (James at the kick-off)", italic=True)
    d.text(3, 16, [
        [("Name: ", {"b": True}), ("James Holloway", {"b": False})],
        [("Role: ", {"b": True}), ("Programme Manager", {"b": False})],
        [("Fun Fact: ", {"b": True}), ("Has run the London Marathon dressed as a hospital bed", {"b": False})],
        [("Personality Test Results: ", {"b": True}), ("Insights: green and blue", {"b": False})],
    ])
    d.text(3, 21, ["Lead a large, multi-team programme from start to close", "Build my coaching skills (ILM level 5 this year)", "Learn more about how social care works with us", "See patients get home earlier"])
    d.text(3, 22, ["Mon to Thu in the office or on wards, Fri at home", "Teams chat for quick questions, email for decisions", "Feedback straight away and in person, please", "15-minute check-in with each lead every week"])
    d.text(3, 23, ["I use dictation software, so emails can be short", "If I go quiet in meetings, ask me what I think", "I collect my children at 5pm on Tuesdays", "I love a good process map, be warned"])
    d.text(4, 17, ["Fewer patients waiting in A&E for a bed", "Better patient experience of discharge", "About 3,000 bed days a year released"], size=11)
    d.text(4, 15, ["Standards the Clinical Group can keep up without us", "Ward teams confident using QI tools", "A model to spread to surgical wards"], size=11)
    d.text(4, 8, ["Patients home before midday, with their medicines", "Ward teams who own their discharge data", "Learning and enjoying the work together"], size=11)
    d.text(4, 28, [[("Success looks like: ", {"b": True}), ("33% of discharges before midday by Sept", {"b": False})]], size=10)
    d.text(4, 27, [[("Success looks like: ", {"b": True}), ("ward teams run board rounds without CITI", {"b": False})]], size=10)
    d.text(4, 26, [[("Success looks like: ", {"b": True}), ("every ward tests a change every fortnight", {"b": False})]], size=10)
    d.text(4, 21, ["Weekly huddle, Mondays 9:30, 30 minutes, hybrid", "Collaboration hours 10am to 4pm; no meetings Friday afternoon", "Teams channel for documents, agendas and the action log", "Agenda out 2 days before; actions out the same day"])
    d.text(4, 9, ["Check in on wellbeing at the start of every huddle", "Raise concerns with James or Priya, or Freedom to Speak Up", "Fortnightly 1:1 with your team lead", "Team lunch on the first Thursday of the month"])
    d.save(dst)


# --------------------------------------------------------------------------
def roles(src, dst):
    d = Deck(src)
    t = d.table(1, 7)
    rows = [
        ("sro", "Chairs the leadership meeting; owns the case for change and the benefits"),
        ("ae", "Chairs escalations with ward leaders; signs off the business case"),
        ("pd", "Leads CITI support, 0.2 WTE"),
        ("pm", "Full time; runs the plan, RAID log and reporting"),
        (["clin", "ops", "nurse"], "One lead per workstream (see organisation chart)"),
        ("pmo", "Owns the Teams site, action log and meeting drum-beat"),
        ("assure", "Gateway review at end of mobilisation (Feb) and before closure (Sep)"),
        (["im1", "im2", "im3", "im4", "analyst"], "Each supports one workstream; Omar supports data"),
        ("supplier", "Agrees morning transport capacity; attends transport group"),
    ]
    for r, (k, comment) in enumerate(rows, 1):
        names = [P[x][0] for x in (k if isinstance(k, list) else [k])]
        fill_cell(t.cell(r, 1), [", ".join(names)], size=7)
        fill_cell(t.cell(r, 3), [comment], size=7)
    d.save(dst)


# --------------------------------------------------------------------------
def org_chart(src, dst):
    d = Deck(src)
    d.delete_slides(1)
    d.para(1, 5, 0, f"Programme Organisation Chart for {NAME}")
    fill_cell(d.table(1, 15).cell(0, 0), [[("Programme Executive Team", {"b": True})], "Dr Amara Okafor (SRO), Tom Reid (Accountable Executive)"])
    t = d.table(1, 13)
    for c, (k, title) in enumerate([("pd", "Programme Director"), ("pm", "Programme Manager"), ("clin", "Clinical Lead"), ("nurse", "Nursing Lead"), ("ops", "Operational Lead")]):
        fill_cell(t.cell(1, c), [P[k][0], title])
    boxes = [(14, 18), (16, 19), (17, 20), (28, 29)]
    for (lead_t, sup_t), ws in zip(boxes, WS):
        t = d.table(1, lead_t)
        fill_cell(t.cell(0, 0), [ws[2]])
        fill_cell(t.cell(2, 0), [ws[3]])
        fill_cell(t.cell(2, 1), [ws[4]])
        fill_cell(t.cell(2, 2), [ws[5]])
        fill_cell(d.table(1, sup_t).cell(1, 0), [ws[6]])
    # the first box has a spare row
    tb = d.table(1, 14)
    if len(tb.rows) > 3:
        for c in range(3):
            fill_cell(tb.cell(3, c), [""])
    for t_id, (wid, name, lead, sup) in zip((21, 24), ENABLING):
        t = d.table(1, t_id)
        fill_cell(t.cell(0, 0), [name])
        fill_cell(t.cell(1, 0), [f"Lead: {lead}   |   CITI support: {sup}"])
    d.delete_shape(1, 32)  # "Blank template to use"
    d.save(dst)


# --------------------------------------------------------------------------
def governance(src, dst):
    d = Deck(src)
    s = 1
    for ident, idx, txt in [
        (9, 0, "Trust Executive Committee"),
        (10, 0, "Medicine Clinical Group Board"),
        (33, 0, "Medicine Clinical Group Executive Committee"),
        (27, 0, NAME),
        (36, 0, "Programme Leadership Meeting (fortnightly)"),
        (43, 0, "Programme weekly huddle"),
        (37, 0, "Ward discharge huddles (daily)"),
        (39, 0, "Workstream user groups"),
        (40, 0, "Transfer of care meeting (twice weekly)"),
        (29, 0, "Medicine senior leadership team (weekly)"),
        (32, 0, "UEC Programme Board (monthly)"),
        (26, 0, "Engagement with other forums"),
        (18, 0, "Medicine Operational Group"),
        (19, 0, "Medicine Clinical Governance"),
        (20, 0, "General Managers' Forum"),
        (21, 0, "Patient Flow Group"),
        (22, 0, "Pharmacy Board"),
        (23, 0, "Patient Experience Committee"),
        (24, 0, "Local authority discharge partnership"),
        (25, 0, "ICB urgent care meeting"),
        (13, 0, "Ward 'Big Room' improvement huddles"),
        (14, 0, "Discharge projects run by single wards"),
    ]:
        d.para(s, ident, idx, txt)
    sh = d.shape(s, 31)
    body = sh.text_frame._txBody
    from officekit import set_para_text
    set_para_text(body, 0, "Urgent and Emergency Care Programme")
    for p in body.findall(q(A, "p")):
        txt = "".join(t.text or "" for t in p.iter(q(A, "t")))
        if txt.startswith("Workstreams"):
            rs = p.findall(q(A, "r"))
            rs[0].find(q(A, "t")).text = "Workstreams: "
            rs[1].find(q(A, "t")).text = "Front door, Same day emergency care, Virtual wards, Discharge flow"
            for r in rs[2:]:
                p.remove(r)
    d.delete_shape(1, 48)
    d.delete_shape(2, 18)
    t = d.table(2, 9)
    fill_cell(t.cell(0, 1), [[("Programme Leadership Meeting", {"b": True})], "(fortnightly, 1 hour)"])
    fill_cell(t.cell(0, 2), [[("Programme Weekly Huddle", {"b": True})], "(Mondays, 30 minutes)"])
    fill_cell(t.cell(0, 3), [[("Workstream User Groups", {"b": True})], "(weekly or fortnightly)"])
    _rows(t, [[None] + r for r in [
        ["Oversees progress against the aim of 33% of discharges before midday. Escalation point for risks and issues; agrees structure, resourcing and budget.",
         "Shares the progress of every workstream (board rounds, TTOs, transport and lounge, complex discharges) against plan, and the links between them.",
         "Gives ward and partner teams a say in each workstream's changes and plans. Each core workstream has its own group."],
        [["Agrees benefits and tracks them through the benefits register", "Agrees roles, resources and the programme plan", "Approves the pharmacy technician business case", "Manages red risks and issues and escalations", "Champions the programme with the Clinical Group"],
         ["Checks delivery against the plan", "Reviews the RAID log and dependencies", "Agrees actions and logs decisions", "Shares learning and celebrates wins"],
         ["Gives feedback on changes before they are tested", "Plans PDSA tests on the wards", "Brings the views of ward staff, patients and partners", "Carries news back to their teams"]],
        [["Leadership report", "Benefits register (monthly)"], ["Highlight reports from each workstream", "RAID log"], ["Ward data and PDSA results"]],
        ["Actions, decisions and escalations", "Actions, decisions and escalations", "Actions, decisions and escalations"],
        [["Chair: Dr Amara Okafor (SRO)", "Tom Reid, Priya Nair, James Holloway", "Workstream leads, Ruth Adeyemi (PMO)"],
         ["Chair: James Holloway", "Workstream leads and CITI support", "Omar Siddiqui (data), Ruth Adeyemi (PMO)"],
         ["Workstream leads, ward champions", "Partners as needed (pharmacy, transport, social work)"]],
    ]], start=1, size=7)
    d.save(dst)


# --------------------------------------------------------------------------
STAKEHOLDERS = [
    ("Ward managers and nurses (8 wards)", "High / High", "Time to take part; staffing on busy days", "Ward huddle, workstream group", "Daily / weekly", "Face to face", "Ward measures, PDSA plans, wins"),
    ("Consultants and junior doctors", "High / Medium", "Board round time; writing TTOs the day before", "Medical staff meeting, Teams post", "Monthly", "Slides, 5-minute slot", "Board round standard, criteria-led discharge"),
    ("Pharmacy team", "High / High", "Workload; technician posts", "Pharmacy team meeting", "Fortnightly", "Face to face", "TTO ready times, rota changes"),
    ("Patient transport provider", "Medium / Low", "Contract terms; vehicle capacity", "Contract meeting", "Monthly", "Meeting, email", "Booking times and volumes"),
    ("Local authority social work", "High / Medium", "Staff time for the daily huddle", "Transfer of care meeting", "Twice weekly", "Teams call", "Complex discharge list, delays"),
    ("Patients and carers", "Low / High", "Knowing when they will go home", "Leaflet, ward posters, patient panel", "Ongoing; panel quarterly", "Print, face to face", "What to expect on the day you go home"),
    ("Clinical Group Executive Committee", "High / Medium", "Bed pressures; money", "Leadership report", "Monthly", "Report", "Progress, risks, decisions needed"),
]


def stakeholders(src, dst):
    d = Deck(src)
    t = d.table(1, 21)
    for (r, c), names in {
        (0, 1): ["Patients and carers", "Ward housekeepers", "Patient panel"],
        (0, 2): ["Ward managers and matrons", "Pharmacy team", "Consultant physicians", "Discharge team"],
        (1, 1): ["Estates (lounge refurbishment)", "Volunteers service"],
        (1, 2): ["Patient transport provider", "Local authority social work", "Clinical Group Executive Committee"],
    }.items():
        body = t.cell(r, c)._tc.txBody
        n = len(body.findall(q(A, "p")))
        set_paras(body, ["Our stakeholders: " + ", ".join(names)], keep=n, template=n - 1, bold=True, size=10, color="002161")
    tt = d.table(2, 7)
    for c, (label, val) in enumerate([("Project Name:", NAME), ("Start date:", long_date(START)), ("Completion date:", long_date(END)), ("Project manager:", P["pm"][0])]):
        fill_cell(tt.cell(0, c), [[(label + " ", {"b": True}), (val, {"b": False})]])
    t = d.table(2, 4)
    add_rows(t, len(STAKEHOLDERS) - 3)
    for r, row in enumerate(STAKEHOLDERS, 1):
        for c, v in enumerate(row):
            fill_cell(t.cell(r, c), [v], size=9)
        t.rows[r].height = Inches(0.5)
    d.save(dst)


# --------------------------------------------------------------------------
PLAN_MONTHS = ["Mar 2026", "Apr 2026", "May 2026", "Jun 2026", "Jul 2026", "Aug 2026"]
# (row, label, start month (0-6), end month) for activity bars
PLAN_BARS = [
    (1, "Test on Wards 5 and 6", 0.0, 1.5), (1, "Spread to all 8 wards", 1.5, 4.0), (1, "Embed and audit", 4.0, 6.0),
    (2, "TTO the day before", 0.0, 2.0), (2, "Technicians on wards", 2.0, 4.5),
    (3, "Morning transport bookings", 0.5, 2.5), (3, "Lounge from 8am", 2.5, 5.0),
    (4, "Daily transfer of care huddle", 0.5, 3.5), (4, "Discharge to assess", 3.5, 6.0),
    (5, "Baseline", 0.0, 1.0), (5, "Weekly SPC charts", 1.0, 6.0),
    (6, "Launch", 0.0, 1.0), (6, "Ward roadshows", 1.0, 3.0), (6, "Handover and closure", 4.5, 6.0),
]
# markers under the bars: (row, month position, label, critical path?)
PLAN_MARKERS = [
    (5, 1.0, "Baseline agreed", True), (1, 1.5, "Wards 5 and 6 tested", True), (2, 2.0, "Technicians in post", True),
    (1, 4.0, "All 8 wards live", True), (6, 6.0, "Handover to Clinical Group", True),
    (3, 2.5, "Lounge opens at 8am", False), (4, 0.5, "First daily huddle", False), (6, 1.0, "Launch event", False),
    (2, 4.5, "Weekend cover agreed", False),
]


def plan_on_a_page(src, dst):
    d = Deck(src)
    s = d.slides[0]
    tshape = d.shape(1, 4)
    t = tshape.table
    fill_cell(t.cell(0, 0), [""])
    for c, m in enumerate(PLAN_MONTHS, 1):
        fill_cell(t.cell(0, c), [m])
    names = [w[2] for w in WS] + ["Data", "PMO & comms"]
    for r, n in enumerate(names, 1):
        fill_cell(t.cell(r, 0), [n])
    d.text(1, 33, "A high level summary of the Discharge Flow Programme from March to August 2026, by workstream, with its milestones. The critical path is the sequence of tasks that must finish on time for the programme to finish on time.")
    x0 = tshape.left + t.columns[0].width
    mw = sum(t.columns[i].width for i in range(1, 7)) / 6
    ys = [tshape.top]
    for row in t.rows:
        ys.append(ys[-1] + row.height)

    bar = d.shape(1, 5)
    dia = d.shape(1, 6)
    tri = d.shape(1, 8)
    lab = d.shape(1, 14)
    keep = {d.shape(1, i)._element for i in (25, 26, 27, 32, 31, 33, 7, 4)}
    proto = {k: copy.deepcopy(v._element) for k, v in (("bar", bar), ("dia", dia), ("tri", tri), ("lab", lab))}
    tree = s.shapes._spTree
    for sh in list(s.shapes):
        if sh._element not in keep:
            tree.remove(sh._element)

    def add(kind, x, y, w=None, h=None, text=None):
        el = copy.deepcopy(proto[kind])
        tree.append(el)
        sh = s.shapes[-1]
        sh.left, sh.top = int(x), int(y)
        if w is not None:
            sh.width = int(w)
        if h is not None:
            sh.height = int(h)
        if text is not None:
            set_paras(sh.text_frame._txBody, [text])
        return sh

    for row, label, a, b in PLAN_BARS:
        add("bar", x0 + a * mw + Emu(20000), ys[row] + Inches(0.06), (b - a) * mw - Emu(40000), Inches(0.24), label)
    pts = []
    for row, pos, label, critical in sorted(PLAN_MARKERS, key=lambda m: m[1]):
        cx = x0 + pos * mw
        top = ys[row] + Inches(0.36)
        if critical:
            add("tri", cx - Inches(0.13), top, Inches(0.26), Inches(0.22))
            pts.append((cx, top + Inches(0.13)))
        else:
            add("dia", cx - Inches(0.11), top, Inches(0.22), Inches(0.22))
        colour = "C00000" if critical else "1C9CC8"
        right = pos > 5
        lb = add("lab", cx - Inches(2.05) if right else cx + Inches(0.15), top - Inches(0.02), Inches(1.9), Inches(0.26))
        set_paras(lb.text_frame._txBody, [label], size=9, color=colour)
        lb.text_frame.word_wrap = False
        lb.text_frame.paragraphs[0].alignment = 3 if right else 1
    for (xa, ya), (xb, yb) in zip(pts, pts[1:]):
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(xa), int(ya), int(xb), int(yb))
        ln.line.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        ln.line.width = Pt(1.25)
        # keep the bars and markers on top of the lines
        tree.remove(ln._element)
        tree.insert(3, ln._element)
    d.save(dst)


# --------------------------------------------------------------------------
def highlight(src, dst):
    d = Deck(src)
    d.text(1, "Title 1", f"{NAME} Highlight Reports")
    d.text(1, "Subtitle 2", long_date(REPORT_DATE))
    d.delete_slides(2)
    s = 2
    d.text(s, "Title 1", "TTOs Workstream Highlight Report")
    d.shape(s, "Title 1").width = Inches(9)
    d.replace(s, 7, "DD MMMM YYYY", long_date(REPORT_DATE))
    d.delete_shape(s, 11)
    t = d.table(s, 4)
    _rows(t, [
        [None, "Write TTOs the day before discharge, put pharmacy technicians on the wards and track TTO ready times on all 8 medical wards."],
        [None, "Arjun Mehta (Pharmacy Technician Lead), Leila Haddad (Lead Pharmacist); CITI support: Hannah Lee"],
        [None, "Pharmacy, 8 medical ward teams, junior doctors, Digital (TTO tracker)"],
    ], start=0, size=9)
    t = d.table(s, 5)
    fill_cell(t.cell(0, 1), ["TTO ready time down from 4.2 to 2.6 hours. Both pharmacy technicians started on 5 May. 71% of TTOs on Wards 5 and 6 now written the day before."], size=9)
    _rows(t, [
        [None, "TTO ready time (hours)", "2.0", "2.6"],
        [None, "TTOs written the day before", "80%", "71%"],
        [None, "Discharges before midday", "33%", "24%"],
    ], start=2, size=9)
    t = d.table(s, 8)
    rows = [
        ("TTO process for the day before discharge", "100%", "B", "10/04", "Agreed by pharmacy and medical staff; on the intranet", "Add to junior doctor induction (August)"),
        ("2 pharmacy technicians on the wards", "90%", "G", "29/05", "Both started 5 May; rota covers 8 wards weekdays", "Agree weekend cover with pharmacy"),
        ("Daily TTO tracker", "60%", "A", "22/05", "Built; ready times on 5 of 8 wards. Data gaps on Wards 2 and 3", "Fix ward codes with Digital (I-002)"),
        ("Spread to all 8 wards", "40%", "R", "05/06", "Wards 1 to 4 behind: junior doctor rota gaps", "Escalate rota gaps to the leadership meeting"),
    ]
    for r, (a, pct, rag, due, prog, nxt) in enumerate(rows, 1):
        fill_cell(t.cell(r, 0), [a], size=9)
        fill_cell(t.cell(r, 1), [pct], size=9, bold=True)
        cell_colour(t.cell(r, 1), rag)
        fill_cell(t.cell(r, 2), [due], size=9)
        fill_cell(t.cell(r, 3), [prog], size=9)
        fill_cell(t.cell(r, 4), [nxt], size=9)
    t = d.table(s, 6)
    fill_cell(t.cell(1, 0), ["D-004: weekend TTO cover to be funded from the technician budget (agreed 28/04)"], size=9)
    fill_cell(t.cell(1, 1), ["Junior doctor rota gaps on Wards 1 to 4 stop TTOs being written the day before. Ask: Clinical Director to raise with the rota team"], size=9)
    fill_cell(t.cell(1, 2), ["R-003 Technicians pulled to dispensary on busy days (amber). I-002 Tracker ward codes wrong on 2 wards (amber)"], size=9)
    fill_cell(t.cell(1, 3), ["A"], size=9, bold=True)
    cell_colour(t.cell(1, 3), "A")
    # Win of the week
    d.text(3, 9, ["“Mr K went home at 10:15 with his medicines explained, two hours earlier than his last stay. His daughter said it was the calmest discharge we've had.”", "Shout out to Arjun and the Ward 6 team"])
    d.save(dst)


# --------------------------------------------------------------------------
def leadership(src, dst):
    d = Deck(src)
    d.text(1, "Title 2", f"{NAME} Leadership Report")
    d.text(1, "Subtitle 3", long_date(LEAD_DATE))
    d.delete_slides(2)
    s = 2
    d.text(s, "Title 1", "Discharge Flow Programme Update")
    d.replace(s, 40, "DD MMMM YYYY", long_date(LEAD_DATE))
    d.delete_shape(s, 3)
    t = d.table(s, 38)
    fill_cell(t.cell(1, 0), [f"{k[0]}: {k[3]} (baseline {k[1]}, target {k[2]})" for k in KPI[:5]], size=8)
    t = d.table(s, 37)
    fill_cell(t.cell(1, 0), ["Ward 6 reached 38% of discharges before midday for 3 weeks running", "Both pharmacy technicians started on 5 May"], size=8)
    t = d.table(s, 31)
    ws_rows = [
        (1, 3, "Board rounds & CLD", "G"), (4, 2, "TTOs & pharmacy", "A"), (6, 2, "Transport, lounge & complex discharges", "G"),
    ]
    for r, span, name, rag in ws_rows:
        fill_cell(t.cell(r, 0), [name], size=8, bold=True)
        cell_colour(t.cell(r, 0), rag)
    lines = [
        ("Morning board round live on 6 of 8 wards", "75%", "G", "05/06", "On track; Wards 1 and 2 start 18 May", "Train Ward 1 and 2 teams"),
        ("Criteria-led discharge policy", "100%", "B", "24/04", "Approved by Clinical Governance on 24 April", "Audit use in June"),
        ("EDD set within 24 hours", "70%", "G", "30/06", "78% of patients, up from 41%", "Add EDD field to board round view (DEP-001)"),
        ("Technicians and TTO the day before", "75%", "A", "05/06", "Wards 1 to 4 behind: junior doctor rota gaps", "Escalation: see risks"),
        ("TTO tracker on all wards", "60%", "A", "22/05", "Data gaps on 2 wards", "Fix ward codes with Digital"),
        ("Lounge open from 8am", "100%", "B", "30/04", "Lounge use up from 12% to 31%", "Promote at ward huddles"),
        ("Daily transfer of care huddle", "80%", "G", "29/05", "Running daily with social work", "Start discharge to assess pilot"),
    ]
    for r, (dl, pct, rag, due, just, nxt) in enumerate(lines, 1):
        fill_cell(t.cell(r, 1), [dl], size=8)
        fill_cell(t.cell(r, 2), [pct], size=8, bold=True)
        cell_colour(t.cell(r, 2), rag)
        fill_cell(t.cell(r, 3), [due], size=8)
        fill_cell(t.cell(r, 4), [just], size=8)
        fill_cell(t.cell(r, 5), [nxt], size=8)
    t = d.table(s, 36)
    for r, (ref, rag, txt) in enumerate([
        ("R-001", "R", "Winter-level bed pressures continue into May, so ward staff miss board round training"),
        ("R-003", "A", "Technicians pulled to dispensary on busy days"),
        ("I-001", "A", "Junior doctor rota gaps on Wards 1 to 4: decision needed on locum cover"),
    ], 2):
        fill_cell(t.cell(r, 0), [ref], size=8)
        fill_cell(t.cell(r, 1), [""], size=8)
        cell_colour(t.cell(r, 1), rag)
        fill_cell(t.cell(r, 2), [txt], size=8)
    d.save(dst)


# --------------------------------------------------------------------------
def closure(src, dst):
    d = Deck(src)
    d.text(1, "Title 1", f"{NAME} Closure Report")
    d.text(1, "Subtitle 2", long_date(END))
    d.text(2, 11, ["Aim: 33% of discharges before midday on 8 medical wards and 0.5 days off average length of stay by September 2026. Closing as planned: CITI support ends and the work is handed to the Clinical Group."], keep=1, size=8)
    d.text(2, 14, ["Discharges before midday up from 16% to 31%", "Length of stay down from 7.4 to 6.9 days", "TTO ready time down from 4.2 to 2.1 hours", "No rise in readmissions"], keep=1, size=8)
    d.text(2, 20, ["Owned by Grace Mensah (General Manager). Weekly SPC charts at the Medicine Operational Group; quarterly Flow Manual audit; ward champions on every ward."], keep=1, size=8)
    t = d.table(2, 13)
    rows = [
        ("Board rounds & CLD", "Board round standard; role cards; CLD policy", "Board round SOP", "Complete", "Joy Adebayo", "Quarterly audit; spread to surgical wards"),
        ("TTOs & pharmacy", "TTO process; 2 technician posts; TTO tracker", "TTO SOP; tracker guide", "Complete", "Leila Haddad", "Agree weekend technician cover"),
        ("Transport & lounge", "Morning bookings; lounge from 8am", "Lounge standard", "Complete", "Kwame Asante", "Review transport contract in 2027"),
        ("Complex discharges", "Daily transfer of care huddle; D2A pilot", "Huddle terms of reference", "Partly complete", "Fiona Gallagher", "Finish D2A evaluation (Nov)"),
        ("Data", "Discharge dashboard with SPC charts", "Dashboard user guide", "Complete", "Omar Siddiqui", "Move to the Trust BI platform"),
        ("PMO & comms", "Patient leaflet; case study; this report", "Comms plan; RAID log", "Complete", "Ruth Adeyemi", "Archive the Teams site in Dec"),
    ]
    _rows(t, [list(r[:2]) + ["Teams: " + r[2]] + list(r[3:]) for r in rows], start=2, size=7)
    t = d.table(3, 7)
    _rows(t, [
        ["Board rounds", "More discharges before midday", "Flow", "16%", "33%", "31%", "Close to target; still improving"],
        ["All", "Shorter length of stay", "Efficiency", "7.4 days", "6.9 days", "6.9 days", "Met; about 3,000 bed days a year"],
        ["TTOs", "Faster TTOs", "Quality", "4.2 hours", "2.0 hours", "2.1 hours", "Close to target"],
        ["Transport & lounge", "More use of the discharge lounge", "Efficiency", "12%", "40%", "38%", "Close to target"],
        ["Board rounds", "EDD set within 24 hours", "Quality", "41%", "90%", "88%", "Close to target"],
        ["All", "No rise in readmissions (balancing)", "Safety", "13.8%", "No rise", "13.5%", "Met"],
    ], start=2, size=8)
    _rows(t, [["Complex discharges", "Fewer patients waiting over 7 days for social care", "Experience", "34 a week", "-", "22 a week", "Not planned; came from the daily huddle"]], start=10, size=8)
    t = d.table(4, 3)
    _rows(t, [
        ["R-007", "02/09/26", "Grace Mensah", "Board rounds", "Gains slip once CITI support ends", "3", "4", "12", "Ward champions; quarterly audit; SPC review monthly", "2", "4", "8", "Handed to Clinical Group"],
        ["R-008", "02/09/26", "Leila Haddad", "TTOs", "No weekend technician cover", "3", "3", "9", "Bid in 2027/28 planning round", "2", "3", "6", "Open"],
    ], start=2, size=7)
    _rows(d.table(4, 8), [["I-006", "09/09/26", "Fiona Gallagher", "Complex discharges", "D2A pilot evaluation not finished", "2", "Finish evaluation with local authority by November", "Handed to discharge team"]], start=2, size=7)
    t = d.table(5, 9)
    _rows(t, [
        ["Ward 6 manager", "“The morning board round is now just how we do it. We don't need CITI in the room any more.”"],
        ["Consultant physician", "“Criteria-led discharge took a while to trust, but it saves me a ward round most days.”"],
        ["Lead pharmacist", "“Having technicians on the wards changed the relationship with ward teams completely.”"],
        ["Patient panel", "“Being told the day before what time I'd go home made a real difference to my family.”"],
        ["Local authority partner", "“The daily huddle means we hear about complex patients days earlier.”"],
    ], start=1, size=9)
    for r in range(1, len(t.rows)):
        for el in list(t.cell(r, 0)._tc.txBody.iter(q(A, "buAutoNum"))):
            el.addprevious(etree.Element(q(A, "buNone")))
            el.getparent().remove(el)
    t = d.table(6, 5)
    _rows(t, [[None] + r for r in [
        ["Mostly: the length of stay aim was met and discharges before midday reached 31% against 33%. Wards 1 to 4 started later because of rota gaps."],
        ["Starting on 2 wards and spreading: other wards saw it working. Weekly SPC charts at ward huddles kept attention on the data."],
        ["Agree the data definitions in mobilisation. We lost 3 weeks fixing ward codes in the TTO tracker."],
        ["Junior doctor rota gaps. Earlier escalation to the Clinical Director would have helped."],
        ["Ward 'Big Rooms' and criteria-led discharge worked well. The patient panel shaped the leaflet."],
        ["Mostly; we were short of analyst time in mobilisation."],
        ["The highlight report template worked well; the RAID log was too long to review weekly, so we filtered to red and amber."],
        ["Share the board round standard with surgical wards first."],
    ]], start=1, size=8)
    d.text(7, 5, [PITCH], keep=1, size=10)
    d.text(7, 10, ["Started on Wards 5 and 6 and spread ward by ward with Plan, Do, Study, Act cycles", "Morning board rounds with criteria-led discharge", "TTOs written the day before, with technicians on the wards", "Lounge open from 8am and morning transport bookings", "Daily transfer of care huddle with social work", "Weekly SPC charts at every ward huddle"], keep=1, size=10)
    d.text(7, 7, ["31% of discharges before midday (from 16%)", "0.5 days off average length of stay", "About 3,000 bed days a year released", "TTOs ready in 2.1 hours (from 4.2)", "No rise in readmissions"], keep=1, size=12, bold=True)
    d.save(dst)
